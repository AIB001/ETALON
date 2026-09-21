"""Agent deliverables: organized structure files, RDKit SDF, Excel, evidence and AF3 jobs."""

import csv
import hashlib
import json
import re
import shutil
from datetime import UTC, datetime
from pathlib import Path

from .._version import __version__
from ..models import utcnow
from .af3 import prepare_af3_jobs
from .compounds import export_compounds, properties
from .review import Curation, integrate_review
from .structures import export_structures


def default_output(targets, mode):
    name = "_".join(t["input"] for t in targets) or "unresolved"
    name = re.sub(r"[^A-Za-z0-9_.-]", "_", name)[:100]
    stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%S_%fZ")
    return Path.cwd() / f"MolQuarry_{name}_{mode}_{stamp}"


def excel_workbook(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        fields = list(dict.fromkeys(key for row in rows for key in row)) if rows else ["status"]
        sheet.append(fields)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24566D")
        for row_number, row in enumerate(rows or [{"status": "No records in this scope"}], 2):
            sheet.append([None] * len(fields))
            for col, field in enumerate(fields, 1):
                value = row.get(field)
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                cell = sheet.cell(row_number, col)
                if isinstance(value, str):
                    # Explicit strings prevent source content from becoming Excel formulas.
                    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
                    cell.value = value[:32767]
                    cell.data_type = "s"
                else:
                    cell.value = value
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for i, field in enumerate(fields, 1):
            sheet.column_dimensions[get_column_letter(i)].width = min(48, max(16, len(field) + 2))
    with Path(path).open("xb") as output:
        book.save(output)


def build_target_bundle(
    quarry,
    dossier_path,
    curation_path,
    output_dir=None,
    *,
    purchasing=True,
    max_purchase_queries=20,
    max_structures=None,
    max_total_bytes=1_000_000_000,
):
    # Optional dependencies are checked before creating an output directory or issuing requests.
    import Bio
    import gemmi
    import numpy
    import openpyxl
    from rdkit import rdBase

    if (
        max_purchase_queries < 0
        or (max_structures is not None and max_structures < 0)
        or max_total_bytes < 1
    ):
        raise ValueError("Query/count budgets must be nonnegative and byte budget positive")

    dossier_path = Path(dossier_path).resolve()
    dossier = json.loads(dossier_path.read_text(encoding="utf-8"))
    curation = (
        Curation.model_validate(curation_path)
        if isinstance(curation_path, dict)
        else Curation.model_validate_json(Path(curation_path).read_text(encoding="utf-8"))
    )
    reviewed = integrate_review(dossier, curation)
    mode = dossier["config"].get("mode", "inhibitor")
    root = Path(output_dir).resolve() if output_dir else default_output(dossier["targets"], mode)
    root.mkdir(parents=True, exist_ok=False)
    evidence = root / "evidence"
    evidence.mkdir()
    # Preserve the exact collection and its relative raw-file locators, without copying caches.
    collection = evidence / "collection"
    collection.mkdir()
    shutil.copyfile(dossier_path, collection / "dossier.json")
    for filename in ["config.json", "summary.json", "manifest.json"]:
        source = dossier_path.parent / filename
        if source.is_file():
            destination = filename
            if filename == "manifest.json" and dossier_path.name != "dossier.json":
                destination = "source_manifest.json"
            shutil.copyfile(source, collection / destination)
    if (dossier_path.parent / "raw").is_dir():
        shutil.copytree(dossier_path.parent / "raw", collection / "raw")
    (evidence / "curation.json").write_text(
        curation.model_dump_json(indent=2) + "\n", encoding="utf-8"
    )
    (evidence / "reviewed.json").write_text(json.dumps(reviewed, indent=2) + "\n", encoding="utf-8")
    structures = export_structures(
        quarry,
        dossier,
        root / "structures" / "experimental",
        max_structures=max_structures,
        max_total_bytes=max_total_bytes,
    )
    compounds = export_compounds(
        quarry,
        dossier,
        reviewed,
        root / "compounds",
        root / "structures" / "experimental",
        purchasing=purchasing,
        max_purchase_queries=max_purchase_queries,
    )
    af3_targets = []
    af3_error = None
    for target in dossier["targets"]:
        target = dict(target)
        if not target.get("sequence"):
            try:
                response = quarry.query("uniprot", "entry", accession=target["accession"])
                target["sequence"] = response.records[0]["sequence"]["value"]
                if (
                    hashlib.sha256(target["sequence"].encode()).hexdigest()
                    != target["sequence_sha256"]
                ):
                    raise ValueError(
                        "Target sequence changed since collection; recollect before modeling"
                    )
                (evidence / f"{target['accession']}_sequence.json").write_text(
                    response.model_dump_json(indent=2) + "\n", encoding="utf-8"
                )
            except Exception as exc:
                af3_error = str(exc)
                break
        af3_targets.append(target)
    try:
        af3 = (
            prepare_af3_jobs(af3_targets, root / "modeling" / "af3" / "jobs")
            if not af3_error
            else {"status": "sequence_unavailable", "reason": af3_error, "models_returned": 0}
        )
    except ValueError as exc:
        af3 = {"status": "invalid_sequence", "reason": str(exc), "models_returned": 0}
    measurements = [
        {"candidate_id": c["candidate_id"], "label": c["label"], **m}
        for c in reviewed["candidates"]
        for m in c["measurements"]
    ]
    molecule_map = {r["molecule_chembl_id"]: r for r in dossier["chembl_details"]["molecule"]}
    document_map = {r["document_chembl_id"]: r for r in dossier["chembl_details"]["document"]}
    database_rows, seen = [], set()
    for item in [*dossier["activities"], *dossier.get("assay_mention_activities", [])]:
        raw = item["record"]
        identity = (raw["activity_id"], item.get("input_target"), item["review_status"])
        if identity in seen:
            continue
        seen.add(identity)
        molecule = molecule_map.get(raw["molecule_chembl_id"], {})
        structure = molecule.get("molecule_structures") or {}
        database_rows.append(
            {
                "activity_id": raw["activity_id"],
                "target_chembl_id": raw["target_chembl_id"],
                "target_name": raw.get("target_pref_name"),
                "compound": raw["molecule_chembl_id"],
                "smiles": structure.get("canonical_smiles") or raw.get("canonical_smiles"),
                "inchikey": structure.get("standard_inchi_key"),
                "measurement": raw.get("standard_type"),
                "relation": raw.get("standard_relation"),
                "value": raw.get("standard_value"),
                "unit": raw.get("standard_units"),
                "assay_id": raw["assay_chembl_id"],
                "assay": raw.get("assay_description"),
                "document": raw.get("document_chembl_id"),
                "pmid": document_map.get(raw.get("document_chembl_id"), {}).get("pubmed_id"),
                "review_status": item["review_status"],
            }
        )
    chemical_rows = []
    for molecule in molecule_map.values():
        descriptor = molecule.get("molecule_structures") or {}
        if descriptor.get("canonical_smiles"):
            try:
                _, props = properties(descriptor["canonical_smiles"])
                chemical_rows.append({"chembl_id": molecule["molecule_chembl_id"], **props})
            except ValueError:
                chemical_rows.append(
                    {"chembl_id": molecule["molecule_chembl_id"], "status": "rdkit_error"}
                )
    papers = [
        {
            k: p.get(k)
            for k in ["id", "source", "pmcid", "doi", "title", "firstPublicationDate", "pubYear"]
        }
        for p in dossier["papers"]
    ]
    tables = root / "tables"
    tables.mkdir()
    sheets = {
        "Candidates": compounds["compounds"],
        "Measurements": measurements,
        "Purchasing": compounds["purchases"],
        "ChemicalProperties": chemical_rows,
        "DatabaseHits": database_rows,
        "Structures": structures["structures"],
        "Alignments": structures["alignments"],
        "Papers": papers,
        "Coverage": dossier["coverage"],
        "Conflicts": reviewed["conflicts"],
        "SDFIndex": compounds["sdf"],
        "AF3": [af3],
    }
    excel_workbook(tables / "compounds.xlsx", sheets)
    with (tables / "structures.csv").open("x", encoding="utf-8", newline="") as file:
        fields = ["pdb_id", "targets", "resolution_angstrom", "method", "status", "pdb_file"]
        writer = csv.DictWriter(file, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(structures["structures"])
    problems = [r for r in structures["structures"] if r["status"] != "pdb_written"]
    result = {
        "schema_version": "1",
        "created_at": utcnow(),
        "output_directory": str(root),
        "mode": mode,
        "targets": [t["accession"] for t in dossier["targets"]],
        "workbook": "tables/compounds.xlsx",
        "candidate_count": len(reviewed["candidates"]),
        "experimental_entries": len(structures["structures"]),
        "pdb_files": sum(r["status"] == "pdb_written" for r in structures["structures"]),
        "aligned_views": sum(
            r["status"] in {"aligned", "reference"} for r in structures["alignments"]
        ),
        "sdf_files": sum(r["status"] == "written" for r in compounds["sdf"]),
        "af3": af3,
        "structure_problems": problems,
        "collection_incomplete_steps": dossier["incomplete_steps"],
        "unreviewed_papers": reviewed["unreviewed_paper_count"],
        "versions": {
            "molquarry": __version__,
            "biopython": Bio.__version__,
            "gemmi": gemmi.__version__,
            "numpy": numpy.__version__,
            "rdkit": rdBase.rdkitVersion,
            "openpyxl": openpyxl.__version__,
        },
        "purchase_scope": compounds["purchase_scope"],
    }
    (root / "run.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    (root / "README.md").write_text(
        "# MolQuarry target deliverables\n\n"
        "Open `tables/compounds.xlsx` for candidates, PMID/SMILES/properties, purchasing evidence, "
        "individual measurements, structures/resolution, alignment and source coverage.\n\n"
        "- `structures/experimental/source_mmcif/`: original coordinates and download manifests.\n"
        "- `structures/experimental/pdb/`: legacy PDB exports; index records conversion limits.\n"
        "- `structures/experimental/aligned/`: target-specific residue-overlap reference groups.\n"
        "- `compounds/source_3d/`: downloaded 3D records, when available.\n"
        "- `compounds/sdf/`: RDKit-prepared poses/conformers, with provenance tags.\n"
        "- `modeling/af3/jobs/`: AF3 inputs and submission state, not completed predictions.\n"
        "- `evidence/`: collection, explicit review, source responses and conflict records.\n"
        "- `run.json` and `manifest.json`: completion counts, limits, versions and checksums.\n\n"
        "Experimental ligand heavy atoms remain fixed during preparation. Generated conformers are "
        "labeled and are not co-crystal poses. A catalog hit does not confirm inventory or price. "
        "Unreviewed rows remain in DatabaseHits; they are not validated inhibitors/agonists. "
        "The reviewed list is limited by the declared search and review coverage.\n",
        encoding="utf-8",
    )
    files = []
    for path in sorted(root.rglob("*")):
        if path.is_file():
            content = path.read_bytes()
            files.append(
                {
                    "path": str(path.relative_to(root)),
                    "bytes": len(content),
                    "sha256": hashlib.sha256(content).hexdigest(),
                }
            )
    (root / "manifest.json").write_text(
        json.dumps({"created_at": utcnow(), "files": files}, indent=2) + "\n", encoding="utf-8"
    )
    return result
