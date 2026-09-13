"""Replaceable local molecule-source plugins with byte-stable provenance."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import shutil
import sqlite3
import stat
import threading
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path, PurePosixPath
from typing import Any, Literal, Self

import pyarrow as pa
import pyarrow.parquet as pq
from pydantic import Field, ValidationError, field_validator, model_validator

from molcascade.config.canonical import canonical_json, canonical_sha256
from molcascade.config.models import StrictFrozenModel
from molcascade.contracts import (
    DECISION_V1,
    RAW_MOLECULE_V1,
    RAW_MOLECULE_V2,
    DataContract,
)
from molcascade.errors import ContractError, PluginError
from molcascade.io.molecules import (
    DELIMITED_FRAMING_VERSION,
    SDF_FRAMING_VERSION,
    FramedRecord,
    SourceSnapshotError,
    extraction_config_sha256,
    iter_delimited_frames,
    iter_sdf_frames,
    read_regular_file_bounded,
    record_provenance,
    snapshot_regular_file,
    source_record_id,
)
from molcascade.io.parquet import validate_parquet_file
from molcascade.plugins.api import PendingOutput, StageContext, StageRequest, StageResponse
from molcascade.plugins.manifest import Cardinality, Determinism, PluginDescriptor, PluginKind

_ColumnRef = str | int
_PROPERTY_HEADER = re.compile(r"^>\s*<([^>]+)>")
_RAW_PATH = Path("datasets/raw/part-00000.parquet")
_DECISION_PATH = Path("datasets/decisions/part-00000.parquet")
_CSV_FIELD_LIMIT_LOCK = threading.Lock()
_DIRECTORY_DESCRIPTOR_TRAVERSAL = (
    os.open in os.supports_dir_fd
    and os.stat in os.supports_dir_fd
    and os.scandir in os.supports_fd
    and bool(getattr(os, "O_DIRECTORY", 0))
    and bool(getattr(os, "O_NOFOLLOW", 0))
)


def _nonblank_path(value: str) -> str:
    if not value or not value.strip() or "\x00" in value:
        raise ValueError("path must be a non-blank string without NUL")
    return value


def _optional_text(value: str | None) -> str | None:
    if value is not None and (not value or "\x00" in value):
        raise ValueError("optional provenance strings must be non-empty and contain no NUL")
    return value


class _LocalMoleculeSourceConfig(StrictFrozenModel):
    schema_version: int = Field(default=1, ge=1, le=1)
    path: str
    source_uri: str | None = None
    generator_id: str | None = None
    batch_id: str | None = None
    batch_size: int = Field(default=65_536, ge=1, le=1_000_000)
    max_record_bytes: int = Field(default=16 * 1024 * 1024, ge=1, le=256 * 1024 * 1024)
    max_buffered_bytes: int = Field(
        default=64 * 1024 * 1024,
        ge=1024,
        le=1024 * 1024 * 1024,
    )

    _validate_path = field_validator("path")(_nonblank_path)
    _validate_optional = field_validator(
        "source_uri",
        "generator_id",
        "batch_id",
    )(_optional_text)

    @model_validator(mode="after")
    def _record_fits_buffer_budget(self) -> Self:
        if self.max_record_bytes > self.max_buffered_bytes:
            raise ValueError("max_record_bytes must not exceed max_buffered_bytes")
        return self


class DelimitedSmilesSourceConfig(_LocalMoleculeSourceConfig):
    """Explicit UTF-8 CSV/TSV/SMI extraction policy; no dialect inference."""

    path: str = "molecules.csv"
    delimiter: str = ","
    has_header: bool = True
    #: Lines to discard before the file proper begins.  Exports that were
    #: assembled for a human reader routinely open with a title or a provenance
    #: banner, and without this the first such line is read as the header and
    #: every configured column is reported missing.  Counted in physical
    #: records, exactly like pandas' ``skiprows``, so ``skip_rows=1`` with
    #: ``has_header`` puts the real header on line 2.
    skip_rows: int = Field(default=0, ge=0, le=1_000_000)
    smiles_column: _ColumnRef = "smiles"
    candidate_id_column: _ColumnRef | None = None
    metadata_columns: tuple[_ColumnRef, ...] = ()
    expected_column_count: int | None = Field(default=None, ge=1, le=10_000)
    quotechar: str = '"'
    escapechar: str | None = None
    doublequote: bool = True
    skipinitialspace: bool = False

    @field_validator("metadata_columns", mode="before")
    @classmethod
    def _metadata_list_to_tuple(cls, value: Any) -> Any:
        return tuple(value) if isinstance(value, list) else value

    @field_validator("delimiter", "quotechar")
    @classmethod
    def _ascii_csv_character(cls, value: str) -> str:
        if len(value) != 1 or not value.isascii() or value in {"\r", "\n", "\x00"}:
            raise ValueError("CSV delimiter and quotechar must be one printable ASCII byte")
        return value

    @field_validator("escapechar")
    @classmethod
    def _ascii_escape_character(cls, value: str | None) -> str | None:
        if value is not None and (
            len(value) != 1
            or not value.isascii()
            or value in {"\r", "\n", "\x00"}
        ):
            raise ValueError("escapechar must be one printable ASCII byte or null")
        return value

    @model_validator(mode="after")
    def _column_reference_mode(self) -> Self:
        references = (self.smiles_column, self.candidate_id_column, *self.metadata_columns)
        present = tuple(value for value in references if value is not None)
        expected_type = str if self.has_header else int
        if any(
            isinstance(value, bool) or not isinstance(value, expected_type)
            for value in present
        ):
            mode = "names" if self.has_header else "zero-based integer indexes"
            raise ValueError(f"column references must use {mode}")
        if not self.has_header and self.expected_column_count is None:
            raise ValueError("headerless input requires expected_column_count")
        if len(self.metadata_columns) != len(set(self.metadata_columns)):
            raise ValueError("metadata_columns contains duplicates")
        if self.delimiter == self.quotechar:
            raise ValueError("delimiter and quotechar must differ")
        if self.escapechar is not None and self.escapechar in {
            self.delimiter,
            self.quotechar,
        }:
            raise ValueError("escapechar must differ from delimiter and quotechar")
        return self


class SDFSourceConfig(_LocalMoleculeSourceConfig):
    """Strict UTF-8 SD-record framing and selected-property extraction."""

    selected_properties: tuple[str, ...] = ()
    candidate_id_property: str | None = None

    @field_validator("selected_properties", mode="before")
    @classmethod
    def _properties_list_to_tuple(cls, value: Any) -> Any:
        return tuple(value) if isinstance(value, list) else value

    @field_validator("selected_properties")
    @classmethod
    def _properties_valid(cls, values: tuple[str, ...]) -> tuple[str, ...]:
        if len(values) != len(set(values)):
            raise ValueError("selected_properties contains duplicates")
        if any(not value or any(char in value for char in "\r\n\x00<>") for value in values):
            raise ValueError("SDF property names must be non-empty plain names")
        return values

    @field_validator("candidate_id_property")
    @classmethod
    def _candidate_property_valid(cls, value: str | None) -> str | None:
        if value is not None and (
            not value or any(char in value for char in "\r\n\x00<>")
        ):
            raise ValueError("candidate_id_property must be a plain property name")
        return value


class RawMoleculeParquetSourceConfig(StrictFrozenModel):
    """All-or-nothing import of an already compliant raw_molecule dataset."""

    schema_version: int = Field(default=1, ge=1, le=1)
    path: str
    batch_size: int = Field(default=65_536, ge=1, le=1_000_000)
    max_buffered_bytes: int = Field(
        default=64 * 1024 * 1024,
        ge=1024,
        le=1024 * 1024 * 1024,
    )

    _validate_path = field_validator("path")(_nonblank_path)


@dataclass(frozen=True, slots=True)
class _SourceEvent:
    raw_row: dict[str, object] | None
    decision_row: dict[str, object] | None


@dataclass(frozen=True, slots=True)
class _WriteSummary:
    input_count: int
    accepted_count: int
    rejected_count: int
    warning_count: int
    decision_count: int
    max_buffered_records: int
    max_buffered_bytes: int


def _ensure_source_request(request: StageRequest) -> None:
    if request.inputs:
        raise PluginError(
            "source plugins do not accept upstream stage inputs",
            code="SOURCE_INPUTS_NOT_ALLOWED",
            context={"input_ports": sorted(request.inputs)},
        )


def _validate_config(model: type[StrictFrozenModel], request: StageRequest) -> Any:
    try:
        return model.model_validate(dict(request.config))
    except ValidationError as error:
        raise PluginError(
            f"invalid source configuration: {error}",
            code="PLUGIN_CONFIG_INVALID",
            context={"error_count": error.error_count()},
        ) from error


def _decision_row(
    *,
    entity_id: str,
    stage_id: str,
    outcome: Literal["REJECT", "WARN"],
    reason_code: str,
    detail: dict[str, object],
) -> dict[str, object]:
    return {
        "entity_id": entity_id,
        "entity_kind": "SOURCE_RECORD",
        "stage_id": stage_id,
        "outcome": outcome,
        "reason_code": reason_code,
        "rule_id": None,
        "detail": canonical_json(detail),
    }


def _flush_rows(
    writer: pq.ParquetWriter,
    rows: list[dict[str, object]],
    schema: pa.Schema,
) -> None:
    if rows:
        writer.write_table(pa.Table.from_pylist(rows, schema=schema))
        rows.clear()


def _estimated_value_bytes(value: object) -> int:
    if value is None:
        return 1
    if isinstance(value, str):
        return len(value.encode("utf-8"))
    if isinstance(value, bytes):
        return len(value)
    if isinstance(value, (bool, int, float)):
        return 8
    return len(repr(value).encode("utf-8"))


def _estimated_event_bytes(event: _SourceEvent) -> int:
    rows = tuple(row for row in (event.raw_row, event.decision_row) if row is not None)
    return sum(
        64 + len(key.encode("utf-8")) + _estimated_value_bytes(value)
        for row in rows
        for key, value in row.items()
    )


def _write_events(
    events: Iterator[_SourceEvent],
    *,
    staging_root: Path,
    batch_size: int,
    max_buffered_bytes: int,
    raw_contract: DataContract = RAW_MOLECULE_V1,
) -> _WriteSummary:
    raw_path = staging_root / _RAW_PATH
    decision_path = staging_root / _DECISION_PATH
    if raw_path.exists() or decision_path.exists():
        raise PluginError(
            "source output path already exists in staging",
            code="PLUGIN_STAGING_NOT_EMPTY",
        )
    raw_path.parent.mkdir(parents=True, exist_ok=True)
    decision_path.parent.mkdir(parents=True, exist_ok=True)
    raw_rows: list[dict[str, object]] = []
    decision_rows: list[dict[str, object]] = []
    input_count = 0
    accepted_count = 0
    rejected_count = 0
    warning_count = 0
    decision_count = 0
    pending = 0
    buffered_bytes = 0
    max_buffered = 0
    observed_max_buffered_bytes = 0
    try:
        with (
            pq.ParquetWriter(raw_path, raw_contract.schema, compression="zstd") as raw_writer,
            pq.ParquetWriter(
                decision_path,
                DECISION_V1.schema,
                compression="zstd",
            ) as decision_writer,
        ):
            for event in events:
                event_bytes = _estimated_event_bytes(event)
                if event_bytes > max_buffered_bytes:
                    raise PluginError(
                        "one expanded source event exceeds max_buffered_bytes",
                        code="SOURCE_BUFFER_BUDGET_EXCEEDED",
                        hint=(
                            "Increase max_buffered_bytes or lower max_record_bytes and "
                            "the number of copied metadata columns."
                        ),
                        context={
                            "estimated_event_bytes": event_bytes,
                            "max_buffered_bytes": max_buffered_bytes,
                        },
                    )
                if pending and (
                    pending >= batch_size
                    or buffered_bytes + event_bytes > max_buffered_bytes
                ):
                    _flush_rows(raw_writer, raw_rows, raw_contract.schema)
                    _flush_rows(decision_writer, decision_rows, DECISION_V1.schema)
                    pending = 0
                    buffered_bytes = 0
                input_count += 1
                pending += 1
                buffered_bytes += event_bytes
                max_buffered = max(max_buffered, pending)
                observed_max_buffered_bytes = max(
                    observed_max_buffered_bytes,
                    buffered_bytes,
                )
                if event.raw_row is not None:
                    accepted_count += 1
                    raw_rows.append(event.raw_row)
                if event.decision_row is not None:
                    decision_count += 1
                    decision_rows.append(event.decision_row)
                    if event.decision_row["outcome"] == "REJECT":
                        rejected_count += 1
                    else:
                        warning_count += 1
                if pending >= batch_size or buffered_bytes >= max_buffered_bytes:
                    _flush_rows(raw_writer, raw_rows, raw_contract.schema)
                    _flush_rows(decision_writer, decision_rows, DECISION_V1.schema)
                    pending = 0
                    buffered_bytes = 0
            if input_count == 0:
                raise PluginError(
                    "source contains no data records",
                    code="SOURCE_EMPTY",
                )
            _flush_rows(raw_writer, raw_rows, raw_contract.schema)
            _flush_rows(decision_writer, decision_rows, DECISION_V1.schema)
        validate_parquet_file(raw_path, expected_schema=raw_contract.schema)
        validate_parquet_file(decision_path, expected_schema=DECISION_V1.schema)
    except BaseException:
        raw_path.unlink(missing_ok=True)
        decision_path.unlink(missing_ok=True)
        raise
    if input_count != accepted_count + rejected_count:
        raw_path.unlink(missing_ok=True)
        decision_path.unlink(missing_ok=True)
        raise PluginError(
            "source record conservation failed",
            code="SOURCE_COUNT_MISMATCH",
            context={
                "input_count": input_count,
                "accepted_count": accepted_count,
                "rejected_count": rejected_count,
            },
        )
    return _WriteSummary(
        input_count,
        accepted_count,
        rejected_count,
        warning_count,
        decision_count,
        max_buffered,
        observed_max_buffered_bytes,
    )


def _stage_response(
    summary: _WriteSummary,
    *,
    metadata: dict[str, Any],
    raw_contract: DataContract = RAW_MOLECULE_V1,
) -> StageResponse:
    counts = {
        "input_record_count": summary.input_count,
        "accepted_record_count": summary.accepted_count,
        "rejected_record_count": summary.rejected_count,
        "warning_record_count": summary.warning_count,
        "decision_row_count": summary.decision_count,
        "max_buffered_records": summary.max_buffered_records,
        "max_buffered_bytes": summary.max_buffered_bytes,
    }
    return StageResponse(
        outputs={
            "primary": PendingOutput(
                contract_id=raw_contract.id,
                file_paths=(_RAW_PATH.as_posix(),),
                metadata={"row_count": summary.accepted_count},
            ),
            "decisions": PendingOutput(
                contract_id=DECISION_V1.id,
                file_paths=(_DECISION_PATH.as_posix(),),
                metadata={"row_count": summary.decision_count},
            ),
        },
        metadata={**metadata, **counts},
    )


def _snapshot(path: str, staging_root: Path, name: str) -> tuple[Path, str, int]:
    snapshot_path = staging_root / name
    try:
        snapshot = snapshot_regular_file(path, snapshot_path)
    except (OSError, SourceSnapshotError) as error:
        raise PluginError(
            f"could not snapshot source file: {error}",
            code="SOURCE_SNAPSHOT_FAILED",
            context={"path": path},
        ) from error
    return snapshot.path, snapshot.sha256, snapshot.size_bytes


def _frame_id(
    frame: FramedRecord,
    *,
    blob_sha256: str,
    source_kind: str,
    framing_version: str,
) -> str:
    return source_record_id(
        source_blob_sha256=blob_sha256,
        source_kind=source_kind,
        framing_version=framing_version,
        record_index=frame.index,
        byte_start=frame.byte_start,
        byte_length=frame.byte_length,
        record_sha256=frame.record_sha256,
    )


def _decode_csv_record(frame: FramedRecord, config: DelimitedSmilesSourceConfig) -> list[str]:
    assert frame.raw_bytes is not None
    text = frame.raw_bytes.decode("utf-8", errors="strict")
    if frame.byte_start == 0 and text.startswith("\ufeff"):
        text = text[1:]
    if "\x00" in text:
        raise ValueError("NUL")
    reader = csv.reader(
        io.StringIO(text, newline=""),
        delimiter=config.delimiter,
        quotechar=config.quotechar,
        escapechar=config.escapechar,
        doublequote=config.doublequote,
        skipinitialspace=config.skipinitialspace,
        strict=True,
    )
    rows = list(reader)
    if len(rows) != 1:
        raise csv.Error(f"framed record decoded into {len(rows)} CSV rows")
    return rows[0]


#: How many header names to quote back when a configured column is absent.
#: Enough to recognise the sheet, short enough to read in a terminal.
_HEADER_PREVIEW = 12


def _header_preview(header: list[str]) -> str:
    shown = ", ".join(repr(name) for name in header[:_HEADER_PREVIEW])
    if len(header) > _HEADER_PREVIEW:
        return f"{shown}, ... ({len(header)} columns)"
    return shown


def _resolve_columns(
    config: DelimitedSmilesSourceConfig,
    header: list[str] | None,
    *,
    kind: str = "CSV",
) -> tuple[int, int | None, tuple[tuple[str, int], ...], int]:
    """Bind configured column names to positions in ``header``.

    ``kind`` names the file being read.  The XLSX reader shares this function,
    and reporting a spreadsheet problem as a "CSV header" problem sends the user
    looking for a format error that is not there -- the workbook parsed fine and
    only the column name was wrong.

    Names are matched exactly first, then case-insensitively.  A file exported
    from ChEMBL calls the column ``SMILES``; one exported from a notebook calls
    it ``smiles``; neither user should have to discover which by trial.  The
    fallback refuses to guess when a header carries two names that differ only
    in case, because there the two spellings are genuinely different columns.
    """

    if header is not None:
        if not header or any(not value for value in header):
            raise PluginError(
                f"{kind} header contains an empty name", code="SOURCE_HEADER_INVALID"
            )
        if len(header) != len(set(header)):
            raise PluginError(
                f"{kind} header contains duplicate names", code="SOURCE_HEADER_INVALID"
            )
        if config.expected_column_count is not None and len(header) != config.expected_column_count:
            raise PluginError(
                f"{kind} header column count does not match expected_column_count",
                code="SOURCE_HEADER_INVALID",
            )
        positions = {name: index for index, name in enumerate(header)}
        folded: dict[str, list[int]] = {}
        for index, name in enumerate(header):
            folded.setdefault(name.casefold(), []).append(index)

        def index_of(reference: _ColumnRef | None) -> int | None:
            if reference is None:
                return None
            assert isinstance(reference, str)
            if reference in positions:
                return positions[reference]
            matches = folded.get(reference.casefold(), [])
            if len(matches) == 1:
                return matches[0]
            if len(matches) > 1:
                spellings = ", ".join(repr(header[index]) for index in matches)
                raise PluginError(
                    f"column {reference!r} matches {len(matches)} {kind} header names that "
                    f"differ only in case ({spellings}); name the one you mean exactly",
                    code="SOURCE_HEADER_INVALID",
                )
            raise PluginError(
                f"configured column is absent from the {kind} header: {reference!r}. "
                f"The header reads: {_header_preview(header)}",
                code="SOURCE_HEADER_INVALID",
            )

        metadata = tuple((str(ref), index_of(ref)) for ref in config.metadata_columns)
        return (
            index_of(config.smiles_column),  # type: ignore[arg-type,return-value]
            index_of(config.candidate_id_column),
            tuple((name, index) for name, index in metadata if index is not None),
            len(header),
        )

    assert isinstance(config.smiles_column, int)
    assert config.expected_column_count is not None
    candidate = config.candidate_id_column
    assert candidate is None or isinstance(candidate, int)
    metadata = tuple((str(index), index) for index in config.metadata_columns)
    references = [config.smiles_column, *(index for _, index in metadata)]
    if candidate is not None:
        references.append(candidate)
    if any(index < 0 or index >= config.expected_column_count for index in references):
        raise PluginError(
            "configured column index is outside expected_column_count",
            code="PLUGIN_CONFIG_INVALID",
        )
    return config.smiles_column, candidate, metadata, config.expected_column_count


class DelimitedSmilesSourcePlugin:
    descriptor = PluginDescriptor(
        id="source.delimited_smiles",
        version="0.1.0",
        kind=PluginKind.SOURCE,
        outputs=(RAW_MOLECULE_V1.id, DECISION_V1.id),
        output_ports={"primary": RAW_MOLECULE_V1.id, "decisions": DECISION_V1.id},
        cardinality=Cardinality.ONE_TO_MANY,
        determinism=Determinism.DETERMINISTIC,
        display_name="Delimited SMILES source",
        description="Explicit UTF-8 CSV/TSV/SMI ingestion with byte-stable provenance.",
    )
    config_model = DelimitedSmilesSourceConfig

    def execute(self, request: StageRequest, context: StageContext) -> StageResponse:
        _ensure_source_request(request)
        config = _validate_config(self.config_model, request)
        context.staging_root.mkdir(parents=True, exist_ok=True)
        snapshot_path, blob_sha, blob_size = _snapshot(
            config.path,
            context.staging_root,
            ".delimited-source.snapshot",
        )
        extraction_policy = {
            "format": "utf8_csv",
            "delimiter": config.delimiter,
            "has_header": config.has_header,
            # Part of the policy, not a convenience: two runs that discard a
            # different number of leading lines read different molecules out of
            # the same bytes, and must not hash to the same extraction.
            "skip_rows": config.skip_rows,
            "smiles_column": config.smiles_column,
            "candidate_id_column": config.candidate_id_column,
            "metadata_columns": list(config.metadata_columns),
            "expected_column_count": config.expected_column_count,
            "quotechar": config.quotechar,
            "escapechar": config.escapechar,
            "doublequote": config.doublequote,
            "skipinitialspace": config.skipinitialspace,
            "bom_policy": "accept_at_file_start",
            "max_record_bytes": config.max_record_bytes,
        }
        extraction_hash = extraction_config_sha256(extraction_policy)
        _CSV_FIELD_LIMIT_LOCK.acquire()
        old_field_limit = csv.field_size_limit()
        try:
            csv.field_size_limit(max(old_field_limit, config.max_record_bytes))
            frames = iter_delimited_frames(
                snapshot_path,
                delimiter=config.delimiter,
                quotechar=config.quotechar,
                escapechar=config.escapechar,
                doublequote=config.doublequote,
                skipinitialspace=config.skipinitialspace,
                max_record_bytes=config.max_record_bytes,
            )
            header: list[str] | None = None
            first_frame: FramedRecord | None = None
            # Discarded frames are never inspected for framing issues.  A banner
            # row is exactly the kind of line that is malformed as CSV -- an
            # unbalanced quote in a title, a stray delimiter in a date -- and
            # refusing to skip it because it is unparseable would defeat the
            # only reason to skip it.
            for position in range(config.skip_rows):
                if next(frames, None) is None:
                    raise PluginError(
                        f"skip_rows is {config.skip_rows} but the file holds only "
                        f"{position} line(s)",
                        code="SOURCE_HEADER_INVALID",
                    )
            if config.has_header:
                first_frame = next(frames, None)
                if first_frame is None or first_frame.issues or first_frame.raw_bytes is None:
                    raise PluginError(
                        "CSV header is missing or malformed",
                        code="SOURCE_HEADER_INVALID",
                    )
                try:
                    header = _decode_csv_record(first_frame, config)
                except (UnicodeDecodeError, csv.Error, ValueError) as error:
                    raise PluginError(
                        f"CSV header is invalid: {error}",
                        code="SOURCE_HEADER_INVALID",
                    ) from error
            smiles_index, candidate_index, metadata_columns, expected_count = _resolve_columns(
                config,
                header,
            )
            # ``data_record_index`` counts molecules, so it has to discount both
            # the banner and the header; ``frame.index`` still counts physical
            # lines, which is what the provenance record needs it to do.
            data_offset = config.skip_rows + (1 if first_frame is not None else 0)

            def events() -> Iterator[_SourceEvent]:
                for original_frame in frames:
                    data_index = original_frame.index - data_offset
                    record_id = _frame_id(
                        original_frame,
                        blob_sha256=blob_sha,
                        source_kind="delimited_smiles",
                        framing_version=DELIMITED_FRAMING_VERSION,
                    )
                    provenance = {
                        **record_provenance(
                            original_frame,
                            source_blob_sha256=blob_sha,
                            source_kind="delimited_smiles",
                            framing_version=DELIMITED_FRAMING_VERSION,
                            extraction_config_hash=extraction_hash,
                        ),
                        "data_record_index": data_index,
                    }
                    frame = original_frame
                    if "record_too_large" in frame.issues:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="SOURCE_RECORD_TOO_LARGE",
                                detail={**provenance, "issues": list(frame.issues)},
                            ),
                        )
                        continue
                    if "unclosed_quote" in frame.issues:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="DELIMITED_FRAMING_ERROR",
                                detail={**provenance, "issues": list(frame.issues)},
                            ),
                        )
                        continue
                    try:
                        row = _decode_csv_record(frame, config)
                    except UnicodeDecodeError as error:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="SOURCE_UTF8_INVALID",
                                detail={**provenance, "byte_error_start": error.start},
                            ),
                        )
                        continue
                    except ValueError as error:
                        reason = (
                            "DELIMITED_NUL_NOT_ALLOWED"
                            if str(error) == "NUL"
                            else "DELIMITED_PARSE_ERROR"
                        )
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code=reason,
                                detail={**provenance, "error": str(error)},
                            ),
                        )
                        continue
                    except csv.Error as error:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="DELIMITED_PARSE_ERROR",
                                detail={**provenance, "error": str(error)},
                            ),
                        )
                        continue
                    if not row:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="DELIMITED_EMPTY_RECORD",
                                detail=provenance,
                            ),
                        )
                        continue
                    if len(row) != expected_count:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="DELIMITED_COLUMN_COUNT_MISMATCH",
                                detail={
                                    **provenance,
                                    "expected_columns": expected_count,
                                    "actual_columns": len(row),
                                },
                            ),
                        )
                        continue
                    smiles = row[smiles_index]
                    if not smiles.strip():
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="DELIMITED_SMILES_EMPTY",
                                detail=provenance,
                            ),
                        )
                        continue
                    selected = {name: row[index] for name, index in metadata_columns}
                    metadata = {
                        **provenance,
                        "extraction": extraction_policy,
                        "columns": selected,
                    }
                    yield _SourceEvent(
                        {
                            "source_record_id": record_id,
                            "source_kind": "delimited_smiles",
                            "source_uri": config.source_uri,
                            "source_index": data_index,
                            "generator_id": config.generator_id,
                            "batch_id": config.batch_id,
                            "source_candidate_id": (
                                row[candidate_index] if candidate_index is not None else None
                            ),
                            "raw_smiles": smiles,
                            "raw_molblock": None,
                            "source_metadata_json": canonical_json(metadata),
                        },
                        None,
                    )

            summary = _write_events(
                events(),
                staging_root=context.staging_root,
                batch_size=config.batch_size,
                max_buffered_bytes=config.max_buffered_bytes,
            )
            return _stage_response(
                summary,
                metadata={
                    "source_kind": "delimited_smiles",
                    "source_blob_sha256": blob_sha,
                    "source_size_bytes": blob_size,
                    "framing_version": DELIMITED_FRAMING_VERSION,
                    "extraction_config_sha256": extraction_hash,
                },
            )
        finally:
            csv.field_size_limit(old_field_limit)
            _CSV_FIELD_LIMIT_LOCK.release()
            snapshot_path.unlink(missing_ok=True)


def _without_line_ending(value: str) -> str:
    if value.endswith("\r\n"):
        return value[:-2]
    if value.endswith(("\r", "\n")):
        return value[:-1]
    return value


def _sdf_molblock_and_properties(
    text: str,
    selected: tuple[str, ...],
) -> tuple[str, dict[str, str], bool]:
    lines = text.splitlines(keepends=True)
    end_index = next(
        (index for index, line in enumerate(lines) if _without_line_ending(line) == "M  END"),
        None,
    )
    if end_index is None:
        return text, {}, False
    molblock = "".join(lines[: end_index + 1])
    wanted = set(selected)
    properties: dict[str, str] = {}
    cursor = end_index + 1
    while cursor < len(lines):
        header = _PROPERTY_HEADER.match(_without_line_ending(lines[cursor]))
        if header is None:
            cursor += 1
            continue
        name = header.group(1)
        cursor += 1
        values: list[str] = []
        while cursor < len(lines):
            value = _without_line_ending(lines[cursor])
            if not value or _PROPERTY_HEADER.match(value):
                break
            values.append(value)
            cursor += 1
        if name in wanted and name not in properties:
            properties[name] = "\n".join(values)
    return molblock, properties, True


class SDFSourcePlugin:
    descriptor = PluginDescriptor(
        id="source.sdf",
        version="0.1.0",
        kind=PluginKind.SOURCE,
        outputs=(RAW_MOLECULE_V1.id, DECISION_V1.id),
        output_ports={"primary": RAW_MOLECULE_V1.id, "decisions": DECISION_V1.id},
        cardinality=Cardinality.ONE_TO_MANY,
        determinism=Determinism.DETERMINISTIC,
        display_name="SDF source",
        description="Byte-framed UTF-8 SDF ingestion preserving original mol blocks.",
    )
    config_model = SDFSourceConfig

    def execute(self, request: StageRequest, context: StageContext) -> StageResponse:
        _ensure_source_request(request)
        config = _validate_config(self.config_model, request)
        context.staging_root.mkdir(parents=True, exist_ok=True)
        snapshot_path, blob_sha, blob_size = _snapshot(
            config.path,
            context.staging_root,
            ".sdf-source.snapshot",
        )
        requested_properties = tuple(
            dict.fromkeys(
                (
                    *config.selected_properties,
                    *(
                        (config.candidate_id_property,)
                        if config.candidate_id_property is not None
                        else ()
                    ),
                )
            )
        )
        extraction_policy = {
            "format": "sdf_utf8",
            "delimiter_policy": "physical_line_column_1_starts_with_$$$$",
            "selected_properties": list(config.selected_properties),
            "candidate_id_property": config.candidate_id_property,
            "bom_policy": "accept_at_file_start",
            "max_record_bytes": config.max_record_bytes,
        }
        extraction_hash = extraction_config_sha256(extraction_policy)
        try:
            frames = iter_sdf_frames(
                snapshot_path,
                max_record_bytes=config.max_record_bytes,
            )

            def events() -> Iterator[_SourceEvent]:
                for frame in frames:
                    record_id = _frame_id(
                        frame,
                        blob_sha256=blob_sha,
                        source_kind="sdf",
                        framing_version=SDF_FRAMING_VERSION,
                    )
                    provenance = record_provenance(
                        frame,
                        source_blob_sha256=blob_sha,
                        source_kind="sdf",
                        framing_version=SDF_FRAMING_VERSION,
                        extraction_config_hash=extraction_hash,
                    )
                    if "record_too_large" in frame.issues:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="SOURCE_RECORD_TOO_LARGE",
                                detail={**provenance, "issues": list(frame.issues)},
                            ),
                        )
                        continue
                    assert frame.payload_bytes is not None
                    try:
                        text = frame.payload_bytes.decode("utf-8", errors="strict")
                    except UnicodeDecodeError as error:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="SOURCE_UTF8_INVALID",
                                detail={**provenance, "byte_error_start": error.start},
                            ),
                        )
                        continue
                    if frame.byte_start == 0 and text.startswith("\ufeff"):
                        text = text[1:]
                    if not text.strip():
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="SDF_EMPTY_RECORD",
                                detail=provenance,
                            ),
                        )
                        continue
                    molblock, properties, has_m_end = _sdf_molblock_and_properties(
                        text,
                        requested_properties,
                    )
                    issues = list(frame.issues)
                    if not has_m_end:
                        issues.append("missing_m_end")
                    metadata = {
                        **provenance,
                        "extraction": extraction_policy,
                        "sdf_properties": {
                            name: properties[name]
                            for name in config.selected_properties
                            if name in properties
                        },
                    }
                    decision = None
                    if issues:
                        decision = _decision_row(
                            entity_id=record_id,
                            stage_id=request.stage_id,
                            outcome="WARN",
                            reason_code="SDF_MALFORMED_RECORD_FORWARDED",
                            detail={**provenance, "issues": sorted(set(issues))},
                        )
                    yield _SourceEvent(
                        {
                            "source_record_id": record_id,
                            "source_kind": "sdf",
                            "source_uri": config.source_uri,
                            "source_index": frame.index,
                            "generator_id": config.generator_id,
                            "batch_id": config.batch_id,
                            "source_candidate_id": (
                                properties.get(config.candidate_id_property)
                                if config.candidate_id_property is not None
                                else None
                            ),
                            "raw_smiles": None,
                            "raw_molblock": molblock,
                            "source_metadata_json": canonical_json(metadata),
                        },
                        decision,
                    )

            summary = _write_events(
                events(),
                staging_root=context.staging_root,
                batch_size=config.batch_size,
                max_buffered_bytes=config.max_buffered_bytes,
            )
            return _stage_response(
                summary,
                metadata={
                    "source_kind": "sdf",
                    "source_blob_sha256": blob_sha,
                    "source_size_bytes": blob_size,
                    "framing_version": SDF_FRAMING_VERSION,
                    "extraction_config_sha256": extraction_hash,
                },
            )
        finally:
            snapshot_path.unlink(missing_ok=True)


def _discover_parquet_files(path: str) -> tuple[Path, tuple[Path, ...]]:
    root = Path(path)
    if root.is_symlink():
        raise PluginError("Parquet source path must not be a symlink", code="SOURCE_PATH_INVALID")
    if root.is_file():
        if root.suffix.lower() != ".parquet":
            raise PluginError(
                "Parquet source file must end in .parquet", code="SOURCE_PATH_INVALID"
            )
        return root.parent, (root,)
    if not root.is_dir():
        raise PluginError("Parquet source path does not exist", code="SOURCE_PATH_INVALID")
    files = tuple(
        sorted(root.rglob("*.parquet"), key=lambda item: item.relative_to(root).as_posix())
    )
    if not files:
        raise PluginError("Parquet source contains no .parquet files", code="SOURCE_EMPTY")
    if any(file.is_symlink() or not file.is_file() for file in files):
        raise PluginError(
            "Parquet source contains a symlink or non-regular file",
            code="SOURCE_PATH_INVALID",
        )
    return root, files


def _source_file_state(path: Path) -> tuple[int, int, int, int, int]:
    value = path.lstat()
    if (
        stat.S_ISLNK(value.st_mode)
        or _has_reparse_point(value)
        or not stat.S_ISREG(value.st_mode)
    ):
        raise SourceSnapshotError(f"source partition is no longer a regular file: {path}")
    return (
        value.st_dev,
        value.st_ino,
        value.st_size,
        value.st_mtime_ns,
        value.st_ctime_ns,
    )


def _finite_json(value: str) -> object:
    def reject_constant(constant: str) -> None:
        raise ValueError(f"non-finite JSON constant {constant}")

    return json.loads(value, parse_constant=reject_constant)


def _validate_parquet_batch(batch: pa.RecordBatch, connection: sqlite3.Connection) -> None:
    ids = batch.column(batch.schema.get_field_index("source_record_id")).to_pylist()
    kinds = batch.column(batch.schema.get_field_index("source_kind")).to_pylist()
    indexes = batch.column(batch.schema.get_field_index("source_index")).to_pylist()
    metadata_index = batch.schema.get_field_index("source_metadata_json")
    metadata_values = (
        batch.column(metadata_index).to_pylist() if metadata_index >= 0 else [None] * batch.num_rows
    )
    for row_number, (record_id, kind, index, metadata) in enumerate(
        zip(ids, kinds, indexes, metadata_values, strict=True)
    ):
        if not isinstance(record_id, str) or not record_id.strip():
            raise PluginError(
                "Parquet source contains an empty source_record_id",
                code="PARQUET_SOURCE_ID_INVALID",
                context={"batch_row": row_number},
            )
        if not isinstance(kind, str) or not kind.strip():
            raise PluginError(
                "Parquet source contains an empty source_kind",
                code="PARQUET_SOURCE_KIND_INVALID",
                context={"source_record_id": record_id},
            )
        if not isinstance(index, int) or index < 0:
            raise PluginError(
                "Parquet source contains a negative source_index",
                code="PARQUET_SOURCE_INDEX_NEGATIVE",
                context={"source_record_id": record_id},
            )
        if metadata is not None:
            try:
                parsed = _finite_json(metadata)
            except (TypeError, ValueError, json.JSONDecodeError) as error:
                raise PluginError(
                    "Parquet source_metadata_json is not finite valid JSON",
                    code="PARQUET_SOURCE_METADATA_JSON_INVALID",
                    context={"source_record_id": record_id},
                ) from error
            if canonical_json(parsed) != metadata:
                raise PluginError(
                    "Parquet source_metadata_json is not canonical JSON",
                    code="PARQUET_SOURCE_METADATA_JSON_INVALID",
                    context={"source_record_id": record_id},
                )
        try:
            connection.execute("INSERT INTO seen_ids (source_record_id) VALUES (?)", (record_id,))
        except sqlite3.IntegrityError as error:
            raise PluginError(
                f"duplicate source_record_id across Parquet partitions: {record_id}",
                code="PARQUET_SOURCE_DUPLICATE_ID",
                context={"source_record_id": record_id},
            ) from error


def _complete_raw_batch(batch: pa.RecordBatch) -> pa.RecordBatch:
    arrays: list[pa.Array] = []
    for field in RAW_MOLECULE_V1.schema:
        index = batch.schema.get_field_index(field.name)
        arrays.append(
            batch.column(index)
            if index >= 0
            else pa.nulls(batch.num_rows, type=field.type)
        )
    return pa.RecordBatch.from_arrays(arrays, schema=RAW_MOLECULE_V1.schema)


class RawMoleculeParquetSourcePlugin:
    descriptor = PluginDescriptor(
        id="source.raw_molecule_parquet",
        version="0.1.0",
        kind=PluginKind.SOURCE,
        outputs=(RAW_MOLECULE_V1.id, DECISION_V1.id),
        output_ports={"primary": RAW_MOLECULE_V1.id, "decisions": DECISION_V1.id},
        cardinality=Cardinality.ONE_TO_MANY,
        determinism=Determinism.DETERMINISTIC,
        display_name="raw_molecule Parquet source",
        description="Strict streamed import of an existing raw_molecule/v1 dataset.",
    )
    config_model = RawMoleculeParquetSourceConfig

    def execute(self, request: StageRequest, context: StageContext) -> StageResponse:
        _ensure_source_request(request)
        config = _validate_config(self.config_model, request)
        source_root, files = _discover_parquet_files(config.path)
        context.staging_root.mkdir(parents=True, exist_ok=True)
        snapshot_root = context.staging_root / ".parquet-source-snapshot"
        database_path = context.staging_root / ".parquet-source-index.sqlite3"
        snapshots: list[tuple[str, Path, str, int]] = []
        connection: sqlite3.Connection | None = None
        try:
            initial_inventory = tuple(
                file.relative_to(source_root).as_posix() for file in files
            )
            initial_states = tuple(_source_file_state(file) for file in files)
            for index, source_file in enumerate(files):
                relative = source_file.relative_to(source_root).as_posix()
                snapshot = snapshot_regular_file(
                    source_file,
                    snapshot_root / f"part-{index:08d}.parquet",
                )
                snapshots.append((relative, snapshot.path, snapshot.sha256, snapshot.size_bytes))
            try:
                _, current_files = _discover_parquet_files(config.path)
                current_inventory = tuple(
                    file.relative_to(source_root).as_posix() for file in current_files
                )
                current_states = tuple(_source_file_state(file) for file in current_files)
            except (OSError, PluginError, SourceSnapshotError) as error:
                raise PluginError(
                    "Parquet source changed during snapshot",
                    code="PARQUET_SOURCE_DATASET_CHANGED",
                ) from error
            if (
                current_inventory != initial_inventory
                or current_states != initial_states
            ):
                raise PluginError(
                    "Parquet source file inventory or identity changed during snapshot",
                    code="PARQUET_SOURCE_DATASET_CHANGED",
                    context={
                        "initial_file_count": len(initial_inventory),
                        "current_file_count": len(current_inventory),
                    },
                )
            dataset_hash = extraction_config_sha256(
                [
                    {"relative_path": relative, "sha256": digest, "size_bytes": size}
                    for relative, _, digest, size in snapshots
                ]
            )
            connection = sqlite3.connect(database_path)
            connection.execute(
                "CREATE TABLE seen_ids (source_record_id TEXT PRIMARY KEY NOT NULL)"
            )

            def events() -> Iterator[_SourceEvent]:
                for relative, snapshot_path, _, _ in snapshots:
                    try:
                        parquet_file = pq.ParquetFile(snapshot_path)
                        RAW_MOLECULE_V1.validate_schema(parquet_file.schema_arrow)
                    except (OSError, pa.ArrowException, ContractError) as error:
                        raise PluginError(
                            f"Parquet partition violates raw_molecule/v1: {relative}: {error}",
                            code="PARQUET_SOURCE_CONTRACT_INVALID",
                            context={"partition": relative},
                        ) from error
                    read_batch_size = min(config.batch_size, 1024)
                    for batch in parquet_file.iter_batches(batch_size=read_batch_size):
                        try:
                            RAW_MOLECULE_V1.validate(batch)
                        except ContractError as error:
                            raise PluginError(
                                f"Parquet batch violates raw_molecule/v1: {relative}: {error}",
                                code="PARQUET_SOURCE_CONTRACT_INVALID",
                                context={"partition": relative},
                            ) from error
                        _validate_parquet_batch(batch, connection)
                        complete = _complete_raw_batch(batch)
                        arrow_budget = max(1, config.max_buffered_bytes // 4)
                        if complete.num_rows and complete.nbytes:
                            conversion_rows = max(
                                1,
                                min(
                                    1024,
                                    complete.num_rows * arrow_budget // complete.nbytes,
                                ),
                            )
                        else:
                            conversion_rows = 1024
                        for offset in range(0, complete.num_rows, conversion_rows):
                            for row in complete.slice(offset, conversion_rows).to_pylist():
                                yield _SourceEvent(row, None)

            summary = _write_events(
                events(),
                staging_root=context.staging_root,
                batch_size=config.batch_size,
                max_buffered_bytes=config.max_buffered_bytes,
            )
            return _stage_response(
                summary,
                metadata={
                    "source_kind": "raw_molecule_parquet",
                    "source_dataset_sha256": dataset_hash,
                    "input_file_count": len(snapshots),
                    "input_size_bytes": sum(item[3] for item in snapshots),
                    "ids_preserved": True,
                },
            )
        except PluginError:
            raise
        except (OSError, SourceSnapshotError, sqlite3.Error, pa.ArrowException) as error:
            raise PluginError(
                f"Parquet source import failed: {error}",
                code="PARQUET_SOURCE_READ_FAILED",
            ) from error
        finally:
            if connection is not None:
                connection.close()
            database_path.unlink(missing_ok=True)
            shutil.rmtree(snapshot_root, ignore_errors=True)


class XlsxSourceConfig(_LocalMoleculeSourceConfig):
    """Explicit worksheet and column policy for streamed XLSX extraction."""

    path: str = "molecules.xlsx"
    sheet_name: str | None = None
    has_header: bool = True
    #: Worksheet rows to discard before the header.  A sheet built for reading
    #: often opens with a merged section banner -- openpyxl surfaces that as one
    #: title cell followed by ``None``s -- and taking it for the header makes an
    #: otherwise ordinary workbook unreadable.  Counted in rows like pandas'
    #: ``skiprows``: the header is row ``skip_rows + 1``, data begins at
    #: ``skip_rows + 2``.
    skip_rows: int = Field(default=0, ge=0, le=1_048_575)
    smiles_column: _ColumnRef = "smiles"
    candidate_id_column: _ColumnRef | None = None
    metadata_columns: tuple[_ColumnRef, ...] = ()
    expected_column_count: int | None = Field(default=None, ge=1, le=10_000)
    max_columns: int = Field(default=10_000, ge=1, le=16_384)

    @field_validator("sheet_name")
    @classmethod
    def _sheet_name_valid(cls, value: str | None) -> str | None:
        if value is not None and (not value or not value.strip() or "\x00" in value):
            raise ValueError("sheet_name must be non-blank and contain no NUL when provided")
        return value

    @field_validator("metadata_columns", mode="before")
    @classmethod
    def _metadata_list_to_tuple(cls, value: Any) -> Any:
        return tuple(value) if isinstance(value, list) else value

    @model_validator(mode="after")
    def _column_reference_mode(self) -> Self:
        references = (self.smiles_column, self.candidate_id_column, *self.metadata_columns)
        present = tuple(value for value in references if value is not None)
        expected_type = str if self.has_header else int
        if any(
            isinstance(value, bool) or not isinstance(value, expected_type)
            for value in present
        ):
            mode = "names" if self.has_header else "zero-based integer indexes"
            raise ValueError(f"column references must use {mode}")
        if not self.has_header and self.expected_column_count is None:
            raise ValueError("headerless input requires expected_column_count")
        if len(self.metadata_columns) != len(set(self.metadata_columns)):
            raise ValueError("metadata_columns contains duplicates")
        return self


def _xlsx_header_error(
    error: PluginError,
    *,
    workbook_sheets: tuple[str, ...],
    selected: str,
    selection: str,
    header_row: int,
) -> PluginError:
    """Re-raise a header failure saying which sheet it was actually read from.

    A workbook assembled for a human reader usually opens with a README tab, and
    the reader defaults to the first visible sheet.  When that happens every
    column looks missing, and the message -- accurate as far as it goes -- sends
    the user hunting for a column that is sitting in the next sheet along.  So
    the sheet is named, and where it was picked rather than asked for, the
    alternatives are listed.
    """

    message = f"{error.message} (worksheet {selected!r}, header row {header_row})"
    hint = error.hint
    if hint is None and selection == "FIRST_VISIBLE" and len(workbook_sheets) > 1:
        others = ", ".join(repr(name) for name in workbook_sheets if name != selected)
        hint = (
            f"No worksheet was requested, so the first visible one was read. "
            f"This workbook also holds: {others}. Choose one with --sheet."
        )
    return PluginError(
        message,
        code=error.code,
        hint=hint,
        context={
            **error.context,
            "sheet_name": selected,
            "sheet_selection": selection,
            "header_row": header_row,
            "worksheets": list(workbook_sheets),
        },
    )


def _xlsx_column_letter(index: int) -> str:
    """``0 -> 'A'``, ``26 -> 'AA'``: a spreadsheet address the user can see.

    Written out rather than taken from ``openpyxl.utils`` because openpyxl is
    imported lazily inside the reader, and a diagnostic message is a poor reason
    to reach for a module that may not be loaded yet.
    """

    letters = ""
    position = index + 1
    while position:
        position, remainder = divmod(position - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def _xlsx_typed_value(value: object) -> dict[str, object]:
    if value is None:
        return {"type": "null", "value": None}
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"type": "time", "value": value.isoformat()}
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, int):
        return {"type": "integer", "value": value}
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("non-finite numeric cell")
        return {"type": "number", "value": value}
    if isinstance(value, str):
        if "\x00" in value:
            raise ValueError("NUL in string cell")
        return {"type": "string", "value": value}
    raise ValueError(f"unsupported XLSX cell value type: {type(value).__name__}")


@dataclass(frozen=True, slots=True)
class _XlsxEmptyCell:
    value: None = None
    data_type: str = "n"


_XLSX_EMPTY_CELL = _XlsxEmptyCell()


def _xlsx_cell_payload(cell: object, column_index: int) -> dict[str, object]:
    return {
        "column_index": column_index,
        "data_type": str(getattr(cell, "data_type", "n")),
        "value": _xlsx_typed_value(getattr(cell, "value", None)),
    }


def _xlsx_record_id(
    *,
    source_blob_sha256: str,
    sheet_name: str,
    excel_row: int,
    row_cells_sha256: str,
) -> str:
    digest = canonical_sha256(
        {
            "scheme": "molcascade.xlsx-row/v1",
            "source_blob_sha256": source_blob_sha256,
            "sheet_name": sheet_name,
            "excel_row": excel_row,
            "row_cells_sha256": row_cells_sha256,
        }
    )
    return f"source-record:sha256:{digest}"


def _xlsx_candidate_id(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        if "\x00" in value:
            raise ValueError("NUL in candidate ID")
        return value
    if isinstance(value, bool):
        raise ValueError("boolean candidate IDs are not allowed")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and math.isfinite(value):
        return format(value, ".17g")
    raise ValueError("candidate ID must be a string or finite number")


class XlsxSourcePlugin:
    """Stream an explicit or first-visible XLSX worksheet without formulas."""

    descriptor = PluginDescriptor(
        id="source.xlsx",
        version="0.1.0",
        kind=PluginKind.SOURCE,
        outputs=(RAW_MOLECULE_V2.id, DECISION_V1.id),
        output_ports={"primary": RAW_MOLECULE_V2.id, "decisions": DECISION_V1.id},
        cardinality=Cardinality.ONE_TO_MANY,
        determinism=Determinism.DETERMINISTIC,
        display_name="XLSX worksheet source",
        description="Read-only XLSX extraction with explicit or first-visible worksheet policy.",
    )
    config_model = XlsxSourceConfig

    def execute(self, request: StageRequest, context: StageContext) -> StageResponse:
        _ensure_source_request(request)
        config = _validate_config(self.config_model, request)
        workbook_suffix = Path(config.path).suffix.lower()
        if workbook_suffix not in {".xlsx", ".xlsm"}:
            raise PluginError(
                "XLSX source path must end in .xlsx or .xlsm",
                code="SOURCE_PATH_INVALID",
            )
        try:
            import openpyxl
        except ImportError as error:  # pragma: no cover - broken core installation
            raise PluginError(
                "XLSX ingestion requires the core openpyxl dependency",
                code="XLSX_DEPENDENCY_MISSING",
                hint="Repair or reinstall MolCascade; openpyxl is a core dependency.",
            ) from error

        context.staging_root.mkdir(parents=True, exist_ok=True)
        snapshot_path, blob_sha, blob_size = _snapshot(
            config.path,
            context.staging_root,
            f".xlsx-source-snapshot{workbook_suffix}",
        )
        value_workbook: object | None = None
        audit_workbook: object | None = None
        try:
            value_workbook = openpyxl.load_workbook(
                snapshot_path,
                read_only=True,
                data_only=True,
                keep_links=False,
                keep_vba=False,
            )
            audit_workbook = openpyxl.load_workbook(
                snapshot_path,
                read_only=True,
                data_only=False,
                keep_links=False,
                keep_vba=False,
            )
            if config.sheet_name is None:
                visible_sheet_names = [
                    sheet.title
                    for sheet in value_workbook.worksheets
                    if sheet.sheet_state == "visible"
                ]
                if not visible_sheet_names:
                    raise PluginError(
                        "XLSX workbook contains no visible worksheet",
                        code="XLSX_VISIBLE_SHEET_NOT_FOUND",
                    )
                selected_sheet_name = visible_sheet_names[0]
                sheet_selection = "FIRST_VISIBLE"
            elif config.sheet_name not in value_workbook.sheetnames:
                raise PluginError(
                    f"worksheet does not exist: {config.sheet_name}",
                    code="XLSX_SHEET_NOT_FOUND",
                    context={"sheet_name": config.sheet_name},
                )
            else:
                selected_sheet_name = config.sheet_name
                sheet_selection = "EXPLICIT"
            value_sheet = value_workbook[selected_sheet_name]
            audit_sheet = audit_workbook[selected_sheet_name]
            value_sheet.reset_dimensions()
            audit_sheet.reset_dimensions()

            extraction_policy = {
                "format": "xlsx",
                "sheet_name": selected_sheet_name,
                "sheet_selection": sheet_selection,
                "has_header": config.has_header,
                # See the delimited source: which rows were discarded is part of
                # what was extracted, not of how it was requested.
                "skip_rows": config.skip_rows,
                "smiles_column": config.smiles_column,
                "candidate_id_column": config.candidate_id_column,
                "metadata_columns": list(config.metadata_columns),
                "expected_column_count": config.expected_column_count,
                "formula_policy": "reject_selected_cells",
                "error_cell_policy": "reject_selected_cells",
                "openpyxl_modes": {"read_only": True, "data_only": True},
                "max_record_bytes": config.max_record_bytes,
            }
            extraction_hash = extraction_config_sha256(extraction_policy)

            header: list[str] | None = None
            header_row = config.skip_rows + 1
            try:
                if config.has_header:
                    audit_header = next(
                        audit_sheet.iter_rows(min_row=header_row, max_row=header_row),
                        (),
                    )
                    if not audit_header or len(audit_header) > config.max_columns:
                        raise PluginError(
                            f"row {header_row} of the worksheet is empty or exceeds "
                            "max_columns, so it cannot be the header",
                            code="SOURCE_HEADER_INVALID",
                        )
                    if any(cell.data_type in {"f", "e"} for cell in audit_header):
                        raise PluginError(
                            f"row {header_row} of the worksheet contains a formula or "
                            "error cell, so it cannot be the header",
                            code="SOURCE_HEADER_INVALID",
                        )
                    header = []
                    for index, cell in enumerate(audit_header):
                        if not isinstance(cell.value, str):
                            # Nearly always a merged banner: openpyxl gives the
                            # title in the first cell and ``None`` for the rest of
                            # the span.  Saying which cell failed and how to move
                            # past it turns an opaque rejection into a one-flag fix.
                            raise PluginError(
                                f"cell {_xlsx_column_letter(index)}{header_row} is "
                                f"not text ({type(cell.value).__name__}), so row "
                                f"{header_row} cannot be the header. If the sheet "
                                "opens with a title or a merged banner, skip it "
                                f"with skip_rows={header_row}.",
                                code="SOURCE_HEADER_INVALID",
                            )
                        header.append(cell.value)
                (
                    smiles_index,
                    candidate_index,
                    metadata_columns,
                    expected_count,
                ) = _resolve_columns(config, header, kind="XLSX")  # type: ignore[arg-type]
            except PluginError as error:
                raise _xlsx_header_error(
                    error,
                    workbook_sheets=tuple(value_workbook.sheetnames),
                    selected=selected_sheet_name,
                    selection=sheet_selection,
                    header_row=header_row,
                ) from error
            if expected_count > config.max_columns:
                raise PluginError(
                    "XLSX configured column count exceeds max_columns",
                    code="PLUGIN_CONFIG_INVALID",
                )
            selected_indexes = tuple(
                sorted(
                    {
                        smiles_index,
                        *(index for _, index in metadata_columns),
                        *((candidate_index,) if candidate_index is not None else ()),
                    }
                )
            )
            data_start = header_row + 1 if config.has_header else header_row

            def events() -> Iterator[_SourceEvent]:
                value_rows = value_sheet.iter_rows(
                    min_row=data_start,
                    min_col=1,
                )
                audit_rows = audit_sheet.iter_rows(
                    min_row=data_start,
                    min_col=1,
                )
                for source_index, (value_row, audit_row) in enumerate(
                    zip(value_rows, audit_rows, strict=True)
                ):
                    excel_row = data_start + source_index
                    try:
                        row_payload = [
                            _xlsx_cell_payload(cell, index)
                            for index, cell in enumerate(audit_row)
                        ]
                        row_payload_json = canonical_json(row_payload)
                    except ValueError as error:
                        row_payload = [
                            {
                                "column_index": index,
                                "data_type": str(getattr(cell, "data_type", "unknown")),
                                "value_type": type(getattr(cell, "value", None)).__name__,
                            }
                            for index, cell in enumerate(audit_row)
                        ]
                        row_payload_json = canonical_json(row_payload)
                        cell_value_error: ValueError | None = error
                    else:
                        cell_value_error = None
                    row_cells_sha = canonical_sha256(row_payload)
                    record_id = _xlsx_record_id(
                        source_blob_sha256=blob_sha,
                        sheet_name=selected_sheet_name,
                        excel_row=excel_row,
                        row_cells_sha256=row_cells_sha,
                    )
                    provenance = {
                        "source_blob_sha256": blob_sha,
                        "source_kind": "xlsx",
                        "sheet_name": selected_sheet_name,
                        "excel_row": excel_row,
                        "source_index": source_index,
                        "row_cells_sha256": row_cells_sha,
                        "extraction_config_sha256": extraction_hash,
                    }
                    actual_column_count = max(len(value_row), len(audit_row))
                    if actual_column_count > config.max_columns:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="XLSX_COLUMN_COUNT_MISMATCH",
                                detail={
                                    **provenance,
                                    "actual_column_count": actual_column_count,
                                    "max_columns": config.max_columns,
                                },
                            ),
                        )
                        continue
                    extra_cells = audit_row[expected_count:]
                    if any(
                        getattr(cell, "value", None) is not None
                        or getattr(cell, "data_type", "n") in {"f", "e"}
                        for cell in extra_cells
                    ):
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="XLSX_COLUMN_COUNT_MISMATCH",
                                detail={
                                    **provenance,
                                    "actual_column_count": actual_column_count,
                                    "expected_column_count": expected_count,
                                },
                            ),
                        )
                        continue
                    value_row = value_row + (_XLSX_EMPTY_CELL,) * (
                        expected_count - len(value_row)
                    )
                    audit_row = audit_row + (_XLSX_EMPTY_CELL,) * (
                        expected_count - len(audit_row)
                    )
                    if len(row_payload_json.encode("utf-8")) > config.max_record_bytes:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="SOURCE_RECORD_TOO_LARGE",
                                detail=provenance,
                            ),
                        )
                        continue
                    selected_audit_cells = [audit_row[index] for index in selected_indexes]
                    if any(cell.data_type == "f" for cell in selected_audit_cells):
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="XLSX_FORMULA_NOT_ALLOWED",
                                detail=provenance,
                            ),
                        )
                        continue
                    if any(cell.data_type == "e" for cell in selected_audit_cells):
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="XLSX_CELL_ERROR",
                                detail=provenance,
                            ),
                        )
                        continue
                    if cell_value_error is not None:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="XLSX_CELL_TYPE_UNSUPPORTED",
                                detail={**provenance, "error": str(cell_value_error)},
                            ),
                        )
                        continue
                    smiles = value_row[smiles_index].value
                    if not isinstance(smiles, str) or not smiles.strip():
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="XLSX_SMILES_EMPTY",
                                detail=provenance,
                            ),
                        )
                        continue
                    try:
                        candidate_id = (
                            _xlsx_candidate_id(value_row[candidate_index].value)
                            if candidate_index is not None
                            else None
                        )
                        metadata_cells = {
                            name: _xlsx_typed_value(value_row[index].value)
                            for name, index in metadata_columns
                        }
                    except ValueError as error:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="XLSX_CELL_TYPE_UNSUPPORTED",
                                detail={**provenance, "error": str(error)},
                            ),
                        )
                        continue
                    metadata = {
                        **provenance,
                        "extraction": extraction_policy,
                        "metadata_cells": metadata_cells,
                    }
                    yield _SourceEvent(
                        {
                            "source_record_id": record_id,
                            "source_kind": "xlsx",
                            "source_uri": config.source_uri,
                            "source_index": source_index,
                            "generator_id": config.generator_id,
                            "batch_id": config.batch_id,
                            "source_candidate_id": candidate_id,
                            "raw_format": "SMILES",
                            "raw_structure": smiles,
                            "source_metadata_json": canonical_json(metadata),
                        },
                        None,
                    )

            summary = _write_events(
                events(),
                staging_root=context.staging_root,
                batch_size=config.batch_size,
                max_buffered_bytes=config.max_buffered_bytes,
                raw_contract=RAW_MOLECULE_V2,
            )
            return _stage_response(
                summary,
                raw_contract=RAW_MOLECULE_V2,
                metadata={
                    "source_kind": "xlsx",
                    "source_blob_sha256": blob_sha,
                    "source_size_bytes": blob_size,
                    "sheet_name": selected_sheet_name,
                    "sheet_selection": sheet_selection,
                    "extraction_config_sha256": extraction_hash,
                    "openpyxl_version": openpyxl.__version__,
                },
            )
        except PluginError:
            raise
        except Exception as error:
            raise PluginError(
                f"XLSX source read failed: {error}",
                code="XLSX_READ_FAILED",
            ) from error
        finally:
            try:
                try:
                    if audit_workbook is not None:
                        audit_workbook.close()
                finally:
                    if value_workbook is not None:
                        value_workbook.close()
            finally:
                snapshot_path.unlink(missing_ok=True)


class Mol2DirectorySourceConfig(_LocalMoleculeSourceConfig):
    """Deterministic, disk-indexed policy for a directory of loose MOL2 files."""

    path: str = "molecules"
    recursive: bool = True
    candidate_id_mode: Literal["RELATIVE_PATH", "NONE"] = "RELATIVE_PATH"
    max_files: int = Field(default=10_000_000, ge=1, le=100_000_000)


def _has_reparse_point(value: os.stat_result) -> bool:
    flag = getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0)
    return bool(flag and getattr(value, "st_file_attributes", 0) & flag)


def _directory_state(path: Path) -> tuple[int, int, int, int]:
    value = path.lstat()
    if (
        stat.S_ISLNK(value.st_mode)
        or _has_reparse_point(value)
        or not stat.S_ISDIR(value.st_mode)
    ):
        raise SourceSnapshotError(f"directory is not regular or became a symlink: {path}")
    return value.st_dev, value.st_ino, value.st_mtime_ns, value.st_ctime_ns


def _directory_descriptor_state(descriptor: int) -> tuple[int, int, int, int]:
    value = os.fstat(descriptor)
    if not stat.S_ISDIR(value.st_mode):
        raise SourceSnapshotError("opened MOL2 directory descriptor is not a directory")
    return value.st_dev, value.st_ino, value.st_mtime_ns, value.st_ctime_ns


def _file_descriptor_state(descriptor: int) -> tuple[int, int, int, int, int]:
    value = os.fstat(descriptor)
    if not stat.S_ISREG(value.st_mode):
        raise SourceSnapshotError("opened MOL2 file descriptor is not a regular file")
    return (
        value.st_dev,
        value.st_ino,
        value.st_size,
        value.st_mtime_ns,
        value.st_ctime_ns,
    )


def _directory_open_flags() -> int:
    return (
        os.O_RDONLY
        | getattr(os, "O_BINARY", 0)
        | getattr(os, "O_CLOEXEC", 0)
        | getattr(os, "O_DIRECTORY", 0)
        | getattr(os, "O_NOFOLLOW", 0)
    )


def _file_open_flags() -> int:
    return (
        os.O_RDONLY
        | getattr(os, "O_BINARY", 0)
        | getattr(os, "O_CLOEXEC", 0)
        | getattr(os, "O_NOFOLLOW", 0)
    )


def _relative_parts(relative_path: str) -> tuple[str, ...]:
    if not relative_path:
        return ()
    candidate = PurePosixPath(relative_path)
    parts = candidate.parts
    if (
        candidate.is_absolute()
        or candidate.as_posix() != relative_path
        or not parts
        or any(part in {"", ".", ".."} for part in parts)
    ):
        raise SourceSnapshotError("manifest contains an unsafe relative path")
    return parts


def _open_relative_directory(root_descriptor: int, relative_path: str) -> int:
    current = os.dup(root_descriptor)
    try:
        for part in _relative_parts(relative_path):
            child = os.open(part, _directory_open_flags(), dir_fd=current)
            os.close(current)
            current = child
        return current
    except BaseException:
        os.close(current)
        raise


def _open_relative_file(root_descriptor: int, relative_path: str) -> int:
    parts = _relative_parts(relative_path)
    if not parts:
        raise SourceSnapshotError("manifest file path is empty")
    parent = os.dup(root_descriptor)
    try:
        for part in parts[:-1]:
            child = os.open(part, _directory_open_flags(), dir_fd=parent)
            os.close(parent)
            parent = child
        return os.open(parts[-1], _file_open_flags(), dir_fd=parent)
    finally:
        os.close(parent)


@dataclass(frozen=True, slots=True)
class _DescriptorRead:
    sha256: str
    size_bytes: int
    data: bytes | None


def _read_descriptor_bounded(descriptor: int, *, max_bytes: int) -> _DescriptorRead:
    before = _file_descriptor_state(descriptor)
    digest = hashlib.sha256()
    size = 0
    buffer = bytearray()
    oversized = False
    with os.fdopen(os.dup(descriptor), "rb") as reader:
        while chunk := reader.read(1024 * 1024):
            digest.update(chunk)
            size += len(chunk)
            if not oversized and size <= max_bytes:
                buffer.extend(chunk)
            else:
                oversized = True
                buffer.clear()
    after = _file_descriptor_state(descriptor)
    if before != after or size != before[2]:
        raise SourceSnapshotError("MOL2 file changed while being read")
    return _DescriptorRead(
        sha256=digest.hexdigest(),
        size_bytes=size,
        data=None if oversized else bytes(buffer),
    )


def _build_mol2_manifest(
    connection: sqlite3.Connection,
    *,
    root: Path,
    root_descriptor: int | None,
    root_state: tuple[int, int, int, int],
    recursive: bool,
    max_files: int,
) -> tuple[int, int]:
    connection.executescript(
        """
        CREATE TABLE directories (
            relative_path TEXT PRIMARY KEY NOT NULL,
            device INTEGER NOT NULL,
            inode INTEGER NOT NULL,
            mtime_ns INTEGER NOT NULL,
            ctime_ns INTEGER NOT NULL,
            processed INTEGER NOT NULL DEFAULT 0
        );
        CREATE TABLE files (
            relative_path TEXT PRIMARY KEY NOT NULL,
            device INTEGER NOT NULL,
            inode INTEGER NOT NULL,
            size_bytes INTEGER NOT NULL,
            mtime_ns INTEGER NOT NULL,
            ctime_ns INTEGER NOT NULL
        );
        CREATE INDEX directories_pending
        ON directories(processed, relative_path COLLATE BINARY);
        """
    )
    connection.execute(
        "INSERT INTO directories VALUES ('', ?, ?, ?, ?, 0)",
        root_state,
    )
    file_count = 0
    ignored_regular_files = 0
    while row := connection.execute(
        "SELECT relative_path, device, inode, mtime_ns, ctime_ns "
        "FROM directories WHERE processed = 0 "
        "ORDER BY relative_path COLLATE BINARY LIMIT 1"
    ).fetchone():
        relative_directory = row[0]
        expected_directory_state = tuple(row[1:])
        directory = (
            root
            if not relative_directory
            else root.joinpath(*PurePosixPath(relative_directory).parts)
        )
        directory_descriptor = (
            _open_relative_directory(root_descriptor, relative_directory)
            if root_descriptor is not None
            else None
        )
        try:
            before = (
                _directory_descriptor_state(directory_descriptor)
                if directory_descriptor is not None
                else _directory_state(directory)
            )
            if before != expected_directory_state:
                raise SourceSnapshotError(
                    f"directory identity changed before enumeration: {relative_directory or '.'}"
                )
            entries_context = os.scandir(
                directory_descriptor if directory_descriptor is not None else directory
            )
            with entries_context as entries:
                for entry in entries:
                    try:
                        entry.name.encode("utf-8", errors="strict")
                    except UnicodeEncodeError as error:
                        raise PluginError(
                            "MOL2 directory contains a filename that is not valid UTF-8",
                            code="MOL2_DIRECTORY_ENTRY_INVALID",
                            context={
                                "reason": "path_not_utf8",
                                "name_bytes_hex": os.fsencode(entry.name).hex()[:256],
                            },
                        ) from error
                    relative_path = (
                        entry.name
                        if not relative_directory
                        else f"{relative_directory}/{entry.name}"
                    )
                    state = (
                        os.stat(
                            entry.name,
                            dir_fd=directory_descriptor,
                            follow_symlinks=False,
                        )
                        if directory_descriptor is not None
                        else entry.stat(follow_symlinks=False)
                    )
                    if stat.S_ISLNK(state.st_mode) or _has_reparse_point(state):
                        raise PluginError(
                            f"MOL2 directory contains a symlink: {relative_path}",
                            code="MOL2_DIRECTORY_ENTRY_INVALID",
                            context={"relative_path": relative_path, "reason": "symlink"},
                        )
                    if stat.S_ISDIR(state.st_mode):
                        if recursive:
                            connection.execute(
                                "INSERT INTO directories VALUES (?, ?, ?, ?, ?, 0)",
                                (
                                    relative_path,
                                    state.st_dev,
                                    state.st_ino,
                                    state.st_mtime_ns,
                                    state.st_ctime_ns,
                                ),
                            )
                        continue
                    if not stat.S_ISREG(state.st_mode):
                        raise PluginError(
                            f"MOL2 directory contains a non-regular entry: {relative_path}",
                            code="MOL2_DIRECTORY_ENTRY_INVALID",
                            context={
                                "relative_path": relative_path,
                                "reason": "not_regular",
                            },
                        )
                    if Path(entry.name).suffix.lower() != ".mol2":
                        ignored_regular_files += 1
                        continue
                    connection.execute(
                        "INSERT INTO files VALUES (?, ?, ?, ?, ?, ?)",
                        (
                            relative_path,
                            state.st_dev,
                            state.st_ino,
                            state.st_size,
                            state.st_mtime_ns,
                            state.st_ctime_ns,
                        ),
                    )
                    file_count += 1
                    if file_count > max_files:
                        raise PluginError(
                            "MOL2 directory exceeds configured max_files",
                            code="MOL2_DIRECTORY_TOO_MANY_FILES",
                            context={"max_files": max_files},
                        )
            after = (
                _directory_descriptor_state(directory_descriptor)
                if directory_descriptor is not None
                else _directory_state(directory)
            )
        finally:
            if directory_descriptor is not None:
                os.close(directory_descriptor)
        if before != after:
            raise PluginError(
                f"MOL2 directory changed while being enumerated: {relative_directory or '.'}",
                code="MOL2_DIRECTORY_CHANGED",
            )
        connection.execute(
            """
            UPDATE directories
            SET device = ?, inode = ?, mtime_ns = ?, ctime_ns = ?, processed = 1
            WHERE relative_path = ?
            """,
            (*after, relative_directory),
        )
    connection.commit()
    return file_count, ignored_regular_files


def _mol2_record_id(file_sha256: str, source_index: int) -> str:
    digest = canonical_sha256(
        {
            "scheme": "molcascade.mol2-file/v1",
            "file_sha256": file_sha256,
            "sorted_source_index": source_index,
        }
    )
    return f"source-record:sha256:{digest}"


class Mol2DirectorySourcePlugin:
    """Ingest millions of loose MOL2 files through a disk-backed manifest."""

    descriptor = PluginDescriptor(
        id="source.mol2_directory",
        version="0.1.0",
        kind=PluginKind.SOURCE,
        outputs=(RAW_MOLECULE_V2.id, DECISION_V1.id),
        output_ports={"primary": RAW_MOLECULE_V2.id, "decisions": DECISION_V1.id},
        cardinality=Cardinality.ONE_TO_MANY,
        determinism=Determinism.DETERMINISTIC,
        display_name="MOL2 directory source",
        description="Disk-indexed deterministic ingestion of loose MOL2 files.",
    )
    config_model = Mol2DirectorySourceConfig

    def execute(self, request: StageRequest, context: StageContext) -> StageResponse:
        _ensure_source_request(request)
        config = _validate_config(self.config_model, request)
        root = Path(config.path)
        try:
            initial_root_state = _directory_state(root)
        except (OSError, SourceSnapshotError) as error:
            raise PluginError(
                "MOL2 source path must be a regular non-symlink directory",
                code="SOURCE_PATH_INVALID",
            ) from error
        context.staging_root.mkdir(parents=True, exist_ok=True)
        database_path = context.staging_root / ".mol2-source-manifest.sqlite3"
        if database_path.exists():
            raise PluginError(
                "MOL2 manifest database already exists in staging",
                code="PLUGIN_STAGING_NOT_EMPTY",
            )
        connection: sqlite3.Connection | None = None
        root_descriptor: int | None = None
        root_resolved: Path | None = None
        traversal_mode = (
            "descriptor_no_follow"
            if _DIRECTORY_DESCRIPTOR_TRAVERSAL
            else "path_reparse_checked_best_effort"
        )
        extraction_policy = {
            "format": "mol2_directory",
            "recursive": config.recursive,
            "extension": ".mol2_case_insensitive",
            "candidate_id_mode": config.candidate_id_mode,
            "path_order": "sqlite_binary_utf8",
            "max_record_bytes": config.max_record_bytes,
        }
        extraction_hash = extraction_config_sha256(extraction_policy)
        manifest_digest = hashlib.sha256()
        try:
            if _DIRECTORY_DESCRIPTOR_TRAVERSAL:
                root_descriptor = os.open(root, _directory_open_flags())
                root_state = _directory_descriptor_state(root_descriptor)
                if root_state != initial_root_state or _directory_state(root) != root_state:
                    raise SourceSnapshotError(
                        "MOL2 root path identity changed while it was being opened"
                    )
            else:
                root_state = _directory_state(root)
                if root_state != initial_root_state:
                    raise SourceSnapshotError(
                        "MOL2 root path identity changed before enumeration"
                    )
                root_resolved = root.resolve(strict=True)
            connection = sqlite3.connect(database_path)
            connection.execute("PRAGMA journal_mode=OFF")
            connection.execute("PRAGMA synchronous=OFF")
            connection.execute("PRAGMA temp_store=FILE")
            file_count, ignored_count = _build_mol2_manifest(
                connection,
                root=root,
                root_descriptor=root_descriptor,
                root_state=root_state,
                recursive=config.recursive,
                max_files=config.max_files,
            )

            def events() -> Iterator[_SourceEvent]:
                cursor = connection.execute(
                    """
                    SELECT relative_path, device, inode, size_bytes, mtime_ns, ctime_ns
                    FROM files ORDER BY relative_path COLLATE BINARY
                    """
                )
                for source_index, row in enumerate(cursor):
                    relative_path = row[0]
                    expected_state = tuple(row[1:])
                    path = root.joinpath(*PurePosixPath(relative_path).parts)
                    file_descriptor: int | None = None
                    try:
                        if root_descriptor is not None:
                            file_descriptor = _open_relative_file(
                                root_descriptor,
                                relative_path,
                            )
                            if _file_descriptor_state(file_descriptor) != expected_state:
                                raise SourceSnapshotError(
                                    "file identity changed after enumeration"
                                )
                            result = _read_descriptor_bounded(
                                file_descriptor,
                                max_bytes=config.max_record_bytes,
                            )
                            if _file_descriptor_state(file_descriptor) != expected_state:
                                raise SourceSnapshotError("file identity changed after read")
                        else:
                            assert root_resolved is not None
                            if not path.resolve(strict=True).is_relative_to(root_resolved):
                                raise SourceSnapshotError(
                                    "file path resolves outside the MOL2 root directory"
                                )
                            if _source_file_state(path) != expected_state:
                                raise SourceSnapshotError(
                                    "file identity changed after enumeration"
                                )
                            result = read_regular_file_bounded(
                                path,
                                max_bytes=config.max_record_bytes,
                            )
                            if _source_file_state(path) != expected_state:
                                raise SourceSnapshotError("file identity changed after read")
                    except (OSError, SourceSnapshotError) as error:
                        raise PluginError(
                            f"MOL2 source changed while being read: {relative_path}",
                            code="MOL2_DIRECTORY_CHANGED",
                            context={"relative_path": relative_path},
                        ) from error
                    finally:
                        if file_descriptor is not None:
                            os.close(file_descriptor)
                    record_id = _mol2_record_id(result.sha256, source_index)
                    manifest_record = canonical_json(
                        {
                            "source_index": source_index,
                            "relative_path": relative_path,
                            "sha256": result.sha256,
                            "size_bytes": result.size_bytes,
                        }
                    ).encode("utf-8")
                    manifest_digest.update(len(manifest_record).to_bytes(8, "big"))
                    manifest_digest.update(manifest_record)
                    provenance = {
                        "source_kind": "mol2_directory",
                        "relative_path": relative_path,
                        "source_index": source_index,
                        "file_sha256": result.sha256,
                        "file_size_bytes": result.size_bytes,
                        "extraction_config_sha256": extraction_hash,
                    }
                    if result.data is None:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="SOURCE_FILE_TOO_LARGE",
                                detail={
                                    **provenance,
                                    "max_record_bytes": config.max_record_bytes,
                                },
                            ),
                        )
                        continue
                    try:
                        text = result.data.decode("utf-8", errors="strict")
                    except UnicodeDecodeError as error:
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="SOURCE_UTF8_INVALID",
                                detail={**provenance, "byte_error_start": error.start},
                            ),
                        )
                        continue
                    if not text.strip():
                        yield _SourceEvent(
                            None,
                            _decision_row(
                                entity_id=record_id,
                                stage_id=request.stage_id,
                                outcome="REJECT",
                                reason_code="MOL2_EMPTY_FILE",
                                detail=provenance,
                            ),
                        )
                        continue
                    metadata = {**provenance, "extraction": extraction_policy}
                    yield _SourceEvent(
                        {
                            "source_record_id": record_id,
                            "source_kind": "mol2_directory",
                            "source_uri": config.source_uri,
                            "source_index": source_index,
                            "generator_id": config.generator_id,
                            "batch_id": config.batch_id,
                            "source_candidate_id": (
                                relative_path
                                if config.candidate_id_mode == "RELATIVE_PATH"
                                else None
                            ),
                            "raw_format": "MOL2",
                            "raw_structure": text,
                            "source_metadata_json": canonical_json(metadata),
                        },
                        None,
                    )

                for directory_row in connection.execute(
                    """
                    SELECT relative_path, device, inode, mtime_ns, ctime_ns
                    FROM directories ORDER BY relative_path COLLATE BINARY
                    """
                ):
                    relative_directory = directory_row[0]
                    directory_descriptor: int | None = None
                    try:
                        if root_descriptor is not None:
                            directory_descriptor = _open_relative_directory(
                                root_descriptor,
                                relative_directory,
                            )
                            current_state = _directory_descriptor_state(
                                directory_descriptor
                            )
                        else:
                            directory = (
                                root
                                if not relative_directory
                                else root.joinpath(
                                    *PurePosixPath(relative_directory).parts
                                )
                            )
                            assert root_resolved is not None
                            if not directory.resolve(strict=True).is_relative_to(
                                root_resolved
                            ):
                                raise SourceSnapshotError(
                                    "directory resolves outside the MOL2 root"
                                )
                            current_state = _directory_state(directory)
                    except (OSError, SourceSnapshotError) as error:
                        raise PluginError(
                            "MOL2 directory changed during ingestion",
                            code="MOL2_DIRECTORY_CHANGED",
                        ) from error
                    finally:
                        if directory_descriptor is not None:
                            os.close(directory_descriptor)
                    if current_state != tuple(directory_row[1:]):
                        raise PluginError(
                            "MOL2 directory changed during ingestion",
                            code="MOL2_DIRECTORY_CHANGED",
                            context={"relative_path": relative_directory or "."},
                        )
                if _directory_state(root) != root_state:
                    raise PluginError(
                        "MOL2 root path changed during ingestion",
                        code="MOL2_DIRECTORY_CHANGED",
                    )

            summary = _write_events(
                events(),
                staging_root=context.staging_root,
                batch_size=config.batch_size,
                max_buffered_bytes=config.max_buffered_bytes,
                raw_contract=RAW_MOLECULE_V2,
            )
            return _stage_response(
                summary,
                raw_contract=RAW_MOLECULE_V2,
                metadata={
                    "source_kind": "mol2_directory",
                    "manifest_file_count": file_count,
                    "ignored_regular_file_count": ignored_count,
                    "manifest_sha256": manifest_digest.hexdigest(),
                    "manifest_hash_scheme": "length_prefixed_canonical_json_utf8/v1",
                    "extraction_config_sha256": extraction_hash,
                    "path_order": "sqlite_binary_utf8",
                    "multiprocessing": False,
                    "directory_traversal_mode": traversal_mode,
                },
            )
        except PluginError:
            raise
        except (OSError, SourceSnapshotError) as error:
            raise PluginError(
                f"MOL2 directory changed while being ingested: {error}",
                code="MOL2_DIRECTORY_CHANGED",
            ) from error
        except sqlite3.Error as error:
            raise PluginError(
                f"MOL2 directory ingestion failed: {error}",
                code="MOL2_DIRECTORY_READ_FAILED",
            ) from error
        finally:
            try:
                try:
                    if connection is not None:
                        connection.close()
                finally:
                    if root_descriptor is not None:
                        os.close(root_descriptor)
            finally:
                database_path.unlink(missing_ok=True)


__all__ = [
    "DelimitedSmilesSourceConfig",
    "DelimitedSmilesSourcePlugin",
    "Mol2DirectorySourceConfig",
    "Mol2DirectorySourcePlugin",
    "RawMoleculeParquetSourceConfig",
    "RawMoleculeParquetSourcePlugin",
    "SDFSourceConfig",
    "SDFSourcePlugin",
    "XlsxSourceConfig",
    "XlsxSourcePlugin",
]
